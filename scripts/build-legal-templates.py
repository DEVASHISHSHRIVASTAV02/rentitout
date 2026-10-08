"""
Build editable RentItOut appliance rental Word drafts.
These are starting templates only. They are not legal advice.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "public" / "legal-templates"
VERSION_DATE = "October 8, 2026"


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_title(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_field_line(doc: Document, label: str, blank: str = "_" * 42) -> None:
    doc.add_paragraph(f"{label}: {blank}")


def add_notice(doc: Document) -> None:
    add_heading(doc, "IMPORTANT NOTICE")
    add_body(
        doc,
        "This document is an editable starting draft provided by RentItOut for convenience only. "
        "It is not legal advice, not a lawyer-reviewed instrument for your state, and not a guarantee "
        "of stamp-duty compliance, registration, insurance cover, or enforceability. RentItOut is only "
        "a connector marketplace and is not a party to this rental. Both parties should fill every blank, "
        "attach supporting records, and obtain advice from a qualified advocate before signing.",
    )
    add_body(doc, f"Draft version date: {VERSION_DATE}")


def add_signature_block(doc: Document, title: str) -> None:
    add_heading(doc, title)
    add_field_line(doc, "Signature")
    add_field_line(doc, "Full name")
    add_field_line(doc, "Date")
    add_field_line(doc, "Place")


def build_short_form() -> Document:
    doc = Document()
    style_document(doc)
    add_title(doc, "APPLIANCE RENTAL AGREEMENT")
    add_subtitle(doc, "Short Form Draft for Movable Appliances")
    add_notice(doc)

    add_heading(doc, "1. AGREEMENT DETAILS")
    add_field_line(doc, "Agreement date")
    add_field_line(doc, "Place of execution (city, state)")
    add_field_line(doc, "Listing ID (if any)")

    add_heading(doc, "2. PARTIES")
    add_body(doc, "Owner / Lessor")
    for label in ["Full name", "Father's / spouse's name (optional)", "Address", "City and PIN", "Phone", "Email", "Government ID type and number"]:
        add_field_line(doc, label)

    add_body(doc, "Renter / Lessee")
    for label in ["Full name", "Father's / spouse's name (optional)", "Address", "City and PIN", "Phone", "Email", "Government ID type and number"]:
        add_field_line(doc, label)

    add_heading(doc, "3. APPLIANCE SCHEDULE")
    for label in [
        "Category",
        "Brand and model",
        "Serial / IMEI / manufacturer number",
        "Colour / capacity / size",
        "Accessories included",
        "Approximate purchase year / invoice reference",
        "Initial condition summary",
        "Handover location",
    ]:
        add_field_line(doc, label)
    add_body(
        doc,
        "Attach photographs of the appliance and accessories as Annexure A, and the completed handover checklist as Annexure B.",
    )

    add_heading(doc, "4. TERM")
    add_field_line(doc, "Start date")
    add_field_line(doc, "End date")
    add_field_line(doc, "Minimum lock-in period (if any)")
    add_field_line(doc, "Renewal terms (if any)")

    add_heading(doc, "5. RENT, DEPOSIT, AND PAYMENT")
    add_field_line(doc, "Monthly rent (INR)")
    add_field_line(doc, "Rent due date each month")
    add_field_line(doc, "Security deposit (INR)")
    add_field_line(doc, "Payment mode and account / UPI details")
    add_field_line(doc, "Late fee after ___ days (INR or %)")
    add_body(
        doc,
        "The Renter shall pay rent in advance for each month unless the Parties write a different schedule below.",
    )
    add_field_line(doc, "Special payment schedule (if any)")

    add_heading(doc, "6. DELIVERY, USE, AND CARE")
    add_bullets(
        doc,
        [
            "The Owner shall deliver, or make available, the appliance in working condition on the start date, unless delayed for a reason written and accepted by both Parties.",
            "The Renter shall use the appliance only for lawful personal or household use at the agreed location.",
            "The Renter shall not sell, pledge, sub-rent, gift, or permanently transfer the appliance.",
            "The Renter shall not relocate the appliance outside the agreed city or premises without the Owner's prior written consent.",
            "The Renter shall report defects, breakdowns, theft, or loss to the Owner promptly, and in any event within 48 hours of discovery.",
        ],
    )

    add_heading(doc, "7. REPAIRS AND DAMAGE")
    add_bullets(
        doc,
        [
            "Ordinary wear and tear shall be the Owner's responsibility, unless the Parties write otherwise.",
            "Damage caused by misuse, negligence, unauthorised repair, or unlawful use shall be the Renter's responsibility.",
            "The Parties shall record the initial condition in Annexure B. Return inspection shall use the same checklist.",
            "Service contact and expected response time: [DETAILS].",
        ],
    )

    add_heading(doc, "8. SECURITY DEPOSIT")
    add_body(
        doc,
        "The Owner may deduct from the security deposit only for amounts that are due and supported by this Agreement, including unpaid rent, agreed late fees, missing accessories listed in the schedule, and repair costs for damage beyond ordinary wear and tear. Any remaining deposit shall be returned within ___ days after return, inspection, and clearance of dues. The Owner shall share a written deduction list with the Renter.",
    )

    add_heading(doc, "9. EARLY TERMINATION AND DEFAULT")
    add_bullets(
        doc,
        [
            "Either Party may end this Agreement by giving ___ days' written notice, subject to any lock-in and early-exit fee written below.",
            "Early-exit fee / lock-in terms: [DETAILS OR 'NONE'].",
            "If the Renter fails to pay rent when due, misuses the appliance, or fails to return it on the end date, the Owner may terminate this Agreement by written notice and take steps permitted by law to recover the appliance and unpaid amounts.",
            "If the Owner fails to deliver a working appliance as agreed, the Renter may terminate by written notice and seek refund of prepaid rent and deposit that is not lawfully due to the Owner.",
        ],
    )

    add_heading(doc, "10. LIABILITY AND INSURANCE")
    add_body(
        doc,
        "Each Party is responsible for its own acts and omissions. The Renter is responsible for loss or third-party claims arising from the Renter's misuse or unlawful use of the appliance. This draft does not create insurance cover. If either Party wants insurance, that Party shall arrange and pay for it separately and write the details here: [INSURANCE DETAILS OR 'NONE'].",
    )

    add_heading(doc, "11. GOVERNING LAW AND DISPUTES")
    add_body(
        doc,
        "This Agreement is governed by the laws of India. Subject to any arbitration clause written below, courts at [CITY, STATE] shall have jurisdiction. Preferred dispute process before court: [discussion / mediation / arbitration details, or 'NONE'].",
    )

    add_heading(doc, "12. NOTICES AND ENTIRE AGREEMENT")
    add_bullets(
        doc,
        [
            "Notices may be sent by hand, registered post, or the email addresses written above.",
            "This Agreement, Annexure A (photos), Annexure B (handover checklist), and any signed amendments form the complete agreement between the Parties for this rental.",
            "If any clause is held invalid, the remaining clauses continue to apply.",
            "Any change must be in writing and signed by both Parties.",
        ],
    )

    add_signature_block(doc, "OWNER / LESSOR")
    add_signature_block(doc, "RENTER / LESSEE")
    add_signature_block(doc, "WITNESS 1")
    add_signature_block(doc, "WITNESS 2")
    return doc


def build_comprehensive() -> Document:
    doc = Document()
    style_document(doc)
    add_title(doc, "APPLIANCE RENTAL AGREEMENT")
    add_subtitle(doc, "Comprehensive Draft for Movable Appliances")
    add_notice(doc)

    add_heading(doc, "1. TITLE AND DATE")
    add_body(
        doc,
        'This Appliance Rental Agreement ("Agreement") is made on [DATE] at [CITY, STATE] between the Owner and the Renter named below.',
    )

    add_heading(doc, "2. PARTIES")
    add_body(doc, "2.1 Owner / Lessor")
    for label in [
        "Full legal name",
        "Residential / business address",
        "City, state, and PIN",
        "Phone",
        "Email",
        "Government ID type and number",
        "Bank / UPI details for rent (optional attach separately)",
    ]:
        add_field_line(doc, label)

    add_body(doc, "2.2 Renter / Lessee")
    for label in [
        "Full legal name",
        "Residential address during rental",
        "City, state, and PIN",
        "Phone",
        "Email",
        "Government ID type and number",
        "Emergency contact name and phone",
    ]:
        add_field_line(doc, label)

    add_body(
        doc,
        'The Owner and the Renter are together called the "Parties" and individually a "Party". RentItOut is not a Party.',
    )

    add_heading(doc, "3. DEFINITIONS")
    add_bullets(
        doc,
        [
            '"Appliance" means the movable equipment described in Schedule 1, including listed accessories.',
            '"Deposit" means the refundable security amount stated in Clause 6.',
            '"Rent" means the periodic charge stated in Clause 6.',
            '"Premises" means the address where the Appliance may be kept and used under this Agreement.',
            '"Business Day" means a day other than Sunday or a public holiday at the place of execution.',
        ],
    )

    add_heading(doc, "4. GRANT OF RENTAL")
    add_body(
        doc,
        "Subject to this Agreement, the Owner rents the Appliance to the Renter for the Term on a non-transferable basis. Ownership of the Appliance remains with the Owner at all times. The Renter receives only a temporary right to use the Appliance.",
    )

    add_heading(doc, "5. TERM, RENEWAL, AND LOCATION")
    add_field_line(doc, "Start date")
    add_field_line(doc, "End date")
    add_field_line(doc, "Lock-in period (if any)")
    add_field_line(doc, "Premises / permitted location")
    add_bullets(
        doc,
        [
            "The Term begins on the Start Date and ends on the End Date unless ended earlier under this Agreement.",
            "Any renewal must be agreed in writing before the End Date, with updated rent if any.",
            "The Renter shall keep the Appliance only at the Premises unless the Owner gives prior written consent to a move.",
        ],
    )

    add_heading(doc, "6. RENT, DEPOSIT, TAXES, AND RECEIPTS")
    add_field_line(doc, "Monthly rent (INR in figures and words)")
    add_field_line(doc, "Due date each calendar month")
    add_field_line(doc, "Security deposit (INR in figures and words)")
    add_field_line(doc, "Payment mode")
    add_field_line(doc, "Late fee after ___ Business Days")
    add_bullets(
        doc,
        [
            "Rent is payable in advance unless the Parties write a different schedule.",
            "The Deposit is not advance rent unless both Parties expressly agree in writing.",
            "Each Party shall bear taxes that the law places on that Party. If GST or other tax applies and is payable by one Party, write the arrangement here: [TAX DETAILS OR 'NONE'].",
            "The Owner shall issue or confirm a payment receipt or statement when reasonably requested.",
            "Any cash payment should be acknowledged in writing on the same day.",
        ],
    )

    add_heading(doc, "7. HANDOVER, CONDITION, AND RECORDS")
    add_bullets(
        doc,
        [
            "On handover, the Parties shall complete Schedule 2 (Handover and Condition Checklist) and may attach photographs as Schedule 3.",
            "The Owner represents that, to the Owner's knowledge, the Appliance is in working condition as recorded at handover, except for defects expressly noted.",
            "The Renter confirms inspection opportunity before accepting the Appliance.",
            "Serial numbers, accessories, and visible defects must be listed. Silence on an accessory means it is not included.",
        ],
    )

    add_heading(doc, "8. RENTER OBLIGATIONS")
    add_numbered(
        doc,
        [
            "Use the Appliance with reasonable care and only for lawful purposes.",
            "Follow ordinary manufacturer guidance for safe use and power supply where applicable.",
            "Not open, tamper with, or authorise unapproved repair of sealed components.",
            "Not create any lien, charge, or third-party claim over the Appliance.",
            "Not use the Appliance for commercial hire to others without written consent.",
            "Allow the Owner reasonable inspection on prior notice of at least ___ hours, except in emergency or suspected misuse.",
            "Notify the Owner of theft, loss, seizure, or material damage within 24 hours of discovery.",
            "Return the Appliance on the End Date, or earlier termination date, with all listed accessories.",
        ],
    )

    add_heading(doc, "9. OWNER OBLIGATIONS")
    add_numbered(
        doc,
        [
            "Deliver or make available the Appliance as agreed, free from undisclosed third-party claims known to the Owner.",
            "Disclose known major defects before handover.",
            "Cooperate for ordinary wear-and-tear repairs that are the Owner's responsibility under Clause 10.",
            "Return the Deposit balance within the timeline in Clause 11 after lawful deductions.",
            "Provide clear payment details and respond to written rental queries within a reasonable time.",
        ],
    )

    add_heading(doc, "10. MAINTENANCE, REPAIRS, AND BREAKDOWNS")
    add_bullets(
        doc,
        [
            "Ordinary wear and tear and manufacturing defects not caused by the Renter are the Owner's responsibility, unless otherwise written.",
            "Misuse, negligence, unauthorised relocation, power misuse, water damage from renter fault, or unauthorised repair is the Renter's responsibility.",
            "If the Appliance fails without Renter fault and remains unusable for more than ___ consecutive days after written notice, the Parties may agree a rent pause, replacement, or termination without early-exit fee for that failure period.",
            "Service contact person and phone: [DETAILS].",
        ],
    )

    add_heading(doc, "11. DEPOSIT ADJUSTMENT AND FINAL SETTLEMENT")
    add_body(
        doc,
        "The Owner may deduct from the Deposit only proven or agreed amounts for unpaid Rent, agreed late fees, missing scheduled accessories, and repair or replacement cost for damage beyond ordinary wear and tear. The Owner shall give the Renter a written settlement note. The balance Deposit shall be returned within ___ days after return and inspection. Any dispute on deductions should first be discussed in writing within ___ days of the settlement note.",
    )

    add_heading(doc, "12. EARLY TERMINATION")
    add_field_line(doc, "Notice period for ordinary early termination")
    add_field_line(doc, "Early-exit fee during lock-in (if any)")
    add_bullets(
        doc,
        [
            "Either Party may terminate by written notice for material breach if the breach is not cured within ___ days of written notice, unless the breach is not reasonably curable.",
            "The Renter may terminate without early-exit fee if the Owner fails to provide a usable Appliance as agreed and does not cure within the period above.",
            "On termination, the Renter shall return the Appliance and the Parties shall complete final inspection and settlement.",
        ],
    )

    add_heading(doc, "13. DEFAULT AND RECOVERY")
    add_body(
        doc,
        "Events of default by the Renter include non-payment of Rent beyond the late period, unauthorised transfer or removal of the Appliance, creating a third-party claim over it, or material breach of use restrictions. On default, the Owner may terminate this Agreement by written notice and pursue recovery of the Appliance and amounts due through processes permitted by law. Self-help recovery must not involve unlawful force or trespass. The Owner should keep written records of notices and attempts to contact the Renter.",
    )

    add_heading(doc, "14. LOSS, THEFT, AND TOTAL DAMAGE")
    add_body(
        doc,
        "If the Appliance is stolen, lost, or totally damaged while in the Renter's custody, the Renter shall notify the Owner promptly and cooperate with any police complaint the Owner reasonably requires. Liability for replacement value or repair shall follow fault and any insurance arrangement written below. Estimated replacement value for reference only: INR [AMOUNT]. Insurance arranged by: [OWNER / RENTER / NONE]. Policy details: [DETAILS OR 'NONE'].",
    )

    add_heading(doc, "15. INDEMNITY AND LIMITATION")
    add_bullets(
        doc,
        [
            "The Renter shall indemnify the Owner against third-party claims arising from the Renter's misuse or unlawful use of the Appliance, except to the extent caused by the Owner's breach or defect known to and concealed by the Owner.",
            "Neither Party excludes liability for fraud, wilful misconduct, or liability that cannot be limited by law.",
            "Subject to the previous sentence, neither Party is liable to the other for indirect or consequential loss such as lost profits, except for unpaid Rent, Deposit issues, or physical loss of the Appliance.",
        ],
    )

    add_heading(doc, "16. FORCE MAJEURE")
    add_body(
        doc,
        "Neither Party is liable for delay caused by events beyond reasonable control, including natural disaster, war, epidemic restrictions, or government action, provided the affected Party gives prompt written notice and resumes performance when reasonably possible. Prolonged force majeure of more than ___ days may allow either Party to terminate without early-exit fee.",
    )

    add_heading(doc, "17. DATA AND PRIVACY BETWEEN THE PARTIES")
    add_body(
        doc,
        "Each Party shall use the other Party's personal details only for performing this rental, identity verification, payment, and lawful dispute handling. RentItOut's own privacy terms apply only to use of the RentItOut website and do not govern this offline Agreement.",
    )

    add_heading(doc, "18. NOTICES")
    add_body(
        doc,
        "Notices under this Agreement must be in writing and are deemed received if delivered by hand, sent by registered post to the address above, or sent by email to the addresses above with delivery/read evidence where available. A Party may update its notice details by written notice to the other Party.",
    )

    add_heading(doc, "19. GOVERNING LAW, JURISDICTION, AND DISPUTE RESOLUTION")
    add_bullets(
        doc,
        [
            "This Agreement is governed by the laws of India.",
            "Subject to arbitration below, courts at [CITY, STATE] shall have exclusive jurisdiction.",
            "Optional arbitration: If both Parties initial here [YES/NO], disputes shall first be referred to a sole arbitrator appointed by mutual consent under the Arbitration and Conciliation Act, 1996. Seat and venue: [CITY]. Language: English / [OTHER].",
            "Before formal proceedings, the Parties should attempt good-faith discussion for at least ___ days after written dispute notice.",
        ],
    )

    add_heading(doc, "20. STAMP DUTY AND FORMALITIES")
    add_body(
        doc,
        "The Parties are jointly responsible for checking whether stamp duty, e-stamp, franking, or any local formality applies in the state of execution or performance, and for completing that formality before or at signing as required. Cost of stamp duty shall be borne by: [OWNER / RENTER / SHARED]. This draft does not itself complete stamp duty or registration.",
    )

    add_heading(doc, "21. GENERAL")
    add_bullets(
        doc,
        [
            "This Agreement may be signed in counterparts, including scanned or electronic copies if both Parties accept them.",
            "If any provision is invalid, the rest remains in effect.",
            "Failure to enforce a right once is not a waiver of that right later.",
            "This Agreement, Schedule 1, Schedule 2, Schedule 3, and signed amendments are the entire agreement for this rental and replace prior oral discussions on the same subject.",
            "Headings are for convenience only.",
        ],
    )

    add_heading(doc, "SCHEDULE 1 — APPLIANCE DETAILS")
    for label in [
        "Category",
        "Brand",
        "Model",
        "Serial / manufacturer number",
        "Colour / capacity / size",
        "Accessories included",
        "Invoice / warranty reference (if any)",
        "Approximate current value (INR)",
        "Known defects disclosed at signing",
    ]:
        add_field_line(doc, label)

    add_heading(doc, "SCHEDULE 2 — HANDOVER CHECKLIST")
    add_body(
        doc,
        "Use the separate Handover and Condition Checklist document, or attach a completed copy here. Both Parties should initial each page of that checklist.",
    )

    add_heading(doc, "SCHEDULE 3 — PHOTOGRAPH LIST")
    add_body(
        doc,
        "List photo file names or attach prints: front, back, side, serial label, accessories, and any pre-existing marks. [LIST OR ATTACH].",
    )

    add_heading(doc, "EXECUTION")
    add_body(
        doc,
        "IN WITNESS WHEREOF the Parties have signed this Agreement on the date first written above.",
    )
    add_signature_block(doc, "OWNER / LESSOR")
    add_signature_block(doc, "RENTER / LESSEE")
    add_signature_block(doc, "WITNESS 1")
    add_signature_block(doc, "WITNESS 2")
    return doc


def build_checklist() -> Document:
    doc = Document()
    style_document(doc)
    add_title(doc, "APPLIANCE HANDOVER AND CONDITION CHECKLIST")
    add_subtitle(doc, "Annexure / Schedule for Movable Appliance Rentals")
    add_notice(doc)

    add_heading(doc, "1. REFERENCE")
    add_field_line(doc, "Linked agreement date")
    add_field_line(doc, "Listing ID (if any)")
    add_field_line(doc, "Owner name")
    add_field_line(doc, "Renter name")
    add_field_line(doc, "Handover date and time")
    add_field_line(doc, "Return inspection date and time")
    add_field_line(doc, "Location of inspection")

    add_heading(doc, "2. APPLIANCE IDENTITY")
    for label in [
        "Category",
        "Brand and model",
        "Serial / manufacturer number",
        "Colour / capacity / size",
        "Power rating / gas type if relevant",
    ]:
        add_field_line(doc, label)

    add_heading(doc, "3. ACCESSORIES AT HANDOVER")
    add_body(doc, "Mark Yes / No / Not applicable, and count where needed.")
    for item in [
        "Remote control",
        "Power cord / adapter",
        "Stand / base / wall brackets",
        "Filters / trays / detachable parts",
        "User manual",
        "Warranty card / invoice copy",
        "Other: [DESCRIBE]",
    ]:
        doc.add_paragraph(f"{item}: Yes / No / N/A    Qty: ____    Condition: ____________________")

    add_heading(doc, "4. WORKING CHECK AT HANDOVER")
    for item in [
        "Powers on",
        "Primary function works (cooling / washing / heating / etc.)",
        "Unusual noise",
        "Leakage / rust / cracks",
        "Display / lights / buttons work",
        "Remote / controls work",
        "Safety cut-off / door lock works if applicable",
    ]:
        doc.add_paragraph(f"{item}: Pass / Fail / N/A    Notes: ____________________")

    add_heading(doc, "5. VISIBLE CONDITION AT HANDOVER")
    add_field_line(doc, "Scratches / dents / stains")
    add_field_line(doc, "Pre-existing defects disclosed")
    add_field_line(doc, "Overall condition rating (Excellent / Good / Fair / Poor)")
    add_body(
        doc,
        "Photo set taken at handover (tick): front [ ] back [ ] sides [ ] serial label [ ] accessories [ ] defects [ ]",
    )
    add_field_line(doc, "Photo file names or album reference")

    add_heading(doc, "6. METER / READING / SETTINGS IF RELEVANT")
    add_field_line(doc, "Any meter reading / cycle count / hours")
    add_field_line(doc, "Installed settings noted")

    add_heading(doc, "7. HANDOVER DECLARATION")
    add_body(
        doc,
        "The Owner confirms the Appliance and listed accessories were handed over as recorded above. The Renter confirms opportunity to inspect and accepts the Appliance in the recorded condition, subject to latent defects not reasonably visible on inspection.",
    )
    add_signature_block(doc, "OWNER SIGN-OFF AT HANDOVER")
    add_signature_block(doc, "RENTER SIGN-OFF AT HANDOVER")

    add_heading(doc, "8. RETURN INSPECTION")
    add_body(doc, "Complete this section only at return.")
    for item in [
        "Appliance returned on time",
        "All listed accessories returned",
        "Working condition compared with handover",
        "New damage beyond ordinary wear and tear",
        "Cleaning condition acceptable",
        "Keys / remotes / manuals returned",
    ]:
        doc.add_paragraph(f"{item}: Yes / No / N/A    Notes: ____________________")

    add_field_line(doc, "New defects or missing items")
    add_field_line(doc, "Estimated deduction discussion (INR), if any")
    add_field_line(doc, "Deposit return timeline agreed")
    add_body(
        doc,
        "Return photo set taken (tick): front [ ] back [ ] sides [ ] serial label [ ] accessories [ ] new defects [ ]",
    )

    add_heading(doc, "9. RETURN DECLARATION")
    add_body(
        doc,
        "The Parties confirm the return inspection notes above. Any Deposit deduction must still follow the signed rental agreement and be supported by a written settlement note.",
    )
    add_signature_block(doc, "OWNER SIGN-OFF AT RETURN")
    add_signature_block(doc, "RENTER SIGN-OFF AT RETURN")
    add_signature_block(doc, "WITNESS (OPTIONAL)")
    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = {
        "appliance-rental-agreement-short-form.docx": build_short_form(),
        "appliance-rental-agreement-comprehensive.docx": build_comprehensive(),
        "appliance-handover-checklist.docx": build_checklist(),
    }
    for name, document in outputs.items():
        path = OUT_DIR / name
        document.save(path)
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
